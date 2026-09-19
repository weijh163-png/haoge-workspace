"""Build a macro-free Excel 365 revenue-import template.

The workbook is deliberately formula driven: users paste source rows into the
input table, while the target-output and audit sheets recalculate through Excel
functions.  Cost-accounting cells are not generated or edited by this tool.
"""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


INPUT_HEADERS = [
    "序号", "档口唯一标志", "分部名称", "项目代码", "项目名称", "项目类型",
    "分组作业名称", "分组作业窗口名称", "签订人", "全额营业额合计", "备注状态", "收入会计", "补充说明",
]
BASE_HEADERS = [
    "序号", "档口唯一标志", "分部名称", "项目代码", "项目名称", "项目类型",
    "核算会计", "分组作业名称", "分组作业窗口名称", "责任人",
]
OVERRIDE_HEADERS = ["项目代码", "项目名称", "分组作业名称", "分组作业窗口名称", "责任人", "营业额覆盖", "覆盖原因"]
OUTPUT_HEADERS = [
    "状态", "数据源行", "标记", "档口唯一标志", "分部名称", "项目代码", "项目名称", "项目类型",
    "分组作业名称", "分组作业窗口名称", "责任人", "当月营业额", "说明",
]

TITLE_FILL = PatternFill("solid", fgColor="1F5F5B")
HEADER_FILL = PatternFill("solid", fgColor="DFF0EC")
WARN_FILL = PatternFill("solid", fgColor="FCE8E6")
NEW_FILL = PatternFill("solid", fgColor="E5F4EE")
MANUAL_FILL = PatternFill("solid", fgColor="FFF4D6")


def find_sheet(workbook, marker: str):
    return next(sheet for sheet in workbook.worksheets if marker in sheet.title)


def clear_generated_sheets(workbook):
    for title in ["使用说明", "营业额数据源", "目标基准", "人工营业额覆盖", "函数目标输出", "营业额导入核验"]:
        if title in workbook.sheetnames:
            del workbook[title]


def add_table(sheet, name: str, ref: str):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True)
    sheet.add_table(table)


def style_headers(sheet, row: int, width_map: list[float]):
    for index, width in enumerate(width_map, start=1):
        cell = sheet.cell(row, index)
        cell.fill = TITLE_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = width


def build_instructions(workbook):
    sheet = workbook.create_sheet("使用说明", 0)
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "营业额自动导入 · 纯函数模板"
    sheet["A1"].font = Font(size=18, bold=True, color="173D3A")
    sheet["A3"] = "使用方式"
    sheet["A3"].font = Font(size=12, bold=True, color="173D3A")
    steps = [
        "1. 在“营业额数据源”页，从第 4 行开始覆盖粘贴当月原稿数据；不要改动第 3 行表头。",
        "2. 如有确认需保留的人工营业额，在“人工营业额覆盖”页维护五项键、营业额和原因。",
        "3. 打开“各项指标明细（最终表）”：营业额列直接引用本文件数据源；新增窗口会自动追加在表格下方。",
        "4. 打开“营业额导入核验”复核状态；点击“数据源行”可跳转到本文件内的源数据。",
        "5. 核验通过后，使用“另存为”保存成当月目标文件。纯 Excel 函数无法自行另存为新文件。",
        "6. 本模板不计算、不修改任何成本数据；仅管理营业额导入和核验。",
    ]
    for row, line in enumerate(steps, start=5):
        sheet.cell(row, 1, line)
        sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A13"] = "环境要求：Microsoft 365 Excel（需支持 LET、XLOOKUP、FILTER、VSTACK、HSTACK、XMATCH）。"
    sheet["A13"].fill = WARN_FILL
    sheet["A13"].alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 118


def monthly_revenue_formula(row: int) -> str:
    """Internal same-workbook revenue lookup, including preserved overrides."""
    return (
        f'=LET(k,D{row}&"|"&E{row}&"|"&H{row}&"|"&I{row}&"|"&J{row},'
        'sk,tblRevenueInput[项目代码]&"|"&tblRevenueInput[项目名称]&"|"&tblRevenueInput[分组作业名称]&"|"&tblRevenueInput[分组作业窗口名称]&"|"&tblRevenueInput[签订人],'
        'ok,tblRevenueOverride[项目代码]&"|"&tblRevenueOverride[项目名称]&"|"&tblRevenueOverride[分组作业名称]&"|"&tblRevenueOverride[分组作业窗口名称]&"|"&tblRevenueOverride[责任人],'
        'ov,IFERROR(XLOOKUP(k,ok,tblRevenueOverride[营业额覆盖],""),""),'
        'IF(ov<>"",ov,IFERROR(XLOOKUP(k,sk,tblRevenueInput[全额营业额合计],""),"")))'
    )


def added_window_formula() -> str:
    """Spill new source windows into the final table without external links."""
    return (
        '=LET('
        'm,tblRevenueInput[档口唯一标志]<>"",'
        'sid,FILTER(tblRevenueInput[档口唯一标志],m,""),'
        'sbranch,FILTER(tblRevenueInput[分部名称],m,""),'
        'scode,FILTER(tblRevenueInput[项目代码],m,""),'
        'sname,FILTER(tblRevenueInput[项目名称],m,""),'
        'stype,FILTER(tblRevenueInput[项目类型],m,""),'
        'sgroup,FILTER(tblRevenueInput[分组作业名称],m,""),'
        'swindow,FILTER(tblRevenueInput[分组作业窗口名称],m,""),'
        'ssigner,FILTER(tblRevenueInput[签订人],m,""),'
        'samount,FILTER(tblRevenueInput[全额营业额合计],m,""),'
        'sk,scode&"|"&sname&"|"&sgroup&"|"&swindow&"|"&ssigner,'
        'sf,scode&"|"&sname&"|"&sgroup&"|"&swindow,'
        'bk,tblTargetBase[项目代码]&"|"&tblTargetBase[项目名称]&"|"&tblTargetBase[分组作业名称]&"|"&tblTargetBase[分组作业窗口名称]&"|"&tblTargetBase[责任人],'
        'bf,tblTargetBase[项目代码]&"|"&tblTargetBase[项目名称]&"|"&tblTargetBase[分组作业名称]&"|"&tblTargetBase[分组作业窗口名称],'
        'newMask,ISNA(XMATCH(sk,bk)), '
        'nid,FILTER(sid,newMask,""),'
        'nstatus,FILTER(IF(ISNUMBER(XMATCH(sf,bf)),"新增（责任人变更）","新增窗口"),newMask,""),'
        'out,HSTACK(IF(nid<>"","新增",""),nid,FILTER(sbranch,newMask,""),FILTER(scode,newMask,""),FILTER(sname,newMask,""),FILTER(stype,newMask,""),IF(nid<>"","",""),FILTER(sgroup,newMask,""),FILTER(swindow,newMask,""),FILTER(ssigner,newMask,""),FILTER(samount,newMask,""),IF(nid<>"","",""),nstatus),'
        'FILTER(out,INDEX(out,,2)<>"","")'
        ')'
    )


def prepare_final_sheet(workbook, template_sheet):
    """Turn the actual final table into an internal-formula revenue target."""
    final_title = "各项指标明细（最终表）"
    template_sheet.title = final_title
    total_row = next(
        row for row in range(3, template_sheet.max_row + 1)
        if template_sheet.cell(row, 2).value in (None, "")
        and isinstance(template_sheet.cell(row, 11).value, str)
        and template_sheet.cell(row, 11).value.startswith("=SUM(K3:")
    )
    for row in range(3, total_row):
        if template_sheet.cell(row, 2).value not in (None, ""):
            template_sheet.cell(row, 11).value = monthly_revenue_formula(row)
            template_sheet.cell(row, 11).number_format = '#,##0.00'

    note_row, spill_row = total_row + 1, total_row + 2
    for row in range(note_row, spill_row + 260):
        for col in range(1, 14):
            template_sheet.cell(row, col).value = None
    template_sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=13)
    note = template_sheet.cell(note_row, 1)
    note.value = "↓ 自动新增窗口：来自“营业额数据源”；责任人变化会单独新增。成本列由成本会计维护。"
    note.fill = HEADER_FILL
    note.font = Font(bold=True, color="173D3A")
    note.alignment = Alignment(vertical="center")
    template_sheet.row_dimensions[note_row].height = 22

    for row in range(spill_row, spill_row + 260):
        for col in range(1, 14):
            source_style = template_sheet.cell(total_row - 1, col)._style
            template_sheet.cell(row, col)._style = copy(source_style)
            template_sheet.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
    template_sheet.cell(spill_row, 1).value = added_window_formula()
    template_sheet.cell(spill_row, 11).number_format = '#,##0.00'
    template_sheet.conditional_formatting.add(
        f"A{spill_row}:M{spill_row + 259}",
        FormulaRule(formula=[f'$A{spill_row}="新增"'], fill=NEW_FILL),
    )
    template_sheet.cell(total_row, 11).value = (
        f'=SUM(K3:K{total_row - 1})+IFERROR(SUM(INDEX(A{spill_row}#,0,11)),0)'
    )
    template_sheet.cell(total_row, 11).number_format = '#,##0.00'
    template_sheet.freeze_panes = "A3"
    return final_title


def build_input_sheet(workbook, source_sheet):
    sheet = workbook.create_sheet("营业额数据源")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "营业额数据源（每月覆盖粘贴 A:M 数据区）"
    sheet["A1"].font = Font(size=15, bold=True, color="173D3A")
    sheet["A2"] = "数据行从第 4 行开始；目标输出以项目代码、项目名称、班组、窗口、责任人五项进行精确匹配。"
    sheet["A2"].alignment = Alignment(wrap_text=True)
    for col, header in enumerate(INPUT_HEADERS, start=1):
        sheet.cell(3, col, header)
    style_headers(sheet, 3, [8, 16, 18, 14, 30, 14, 18, 34, 14, 18, 14, 14, 30])

    last_row = source_sheet.max_row
    for source_row in range(4, last_row + 1):
        for col in range(1, 14):
            source_cell = source_sheet.cell(source_row, col)
            target = sheet.cell(source_row, col, source_cell.value)
            target.number_format = source_cell.number_format
            target.alignment = Alignment(vertical="top", wrap_text=True)
    add_table(sheet, "tblRevenueInput", f"A3:M{last_row}")
    sheet.freeze_panes = "A4"
    return sheet


def build_base_sheet(workbook, template_sheet):
    sheet = workbook.create_sheet("目标基准")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "目标基准（来自成本明细模板；作为收入侧匹配基准）"
    sheet["A1"].font = Font(size=15, bold=True, color="173D3A")
    for col, header in enumerate(BASE_HEADERS, start=1):
        sheet.cell(3, col, header)
    style_headers(sheet, 3, [8, 16, 18, 14, 30, 14, 14, 18, 34, 14])
    records = []
    for row in range(3, template_sheet.max_row + 1):
        if template_sheet.cell(row, 2).value in (None, ""):
            continue
        records.append([template_sheet.cell(row, col).value for col in range(1, 11)])
    for output_row, values in enumerate(records, start=4):
        for col, value in enumerate(values, start=1):
            sheet.cell(output_row, col, value)
            sheet.cell(output_row, col).alignment = Alignment(vertical="top", wrap_text=True)
    add_table(sheet, "tblTargetBase", f"A3:J{3 + len(records)}")
    sheet.freeze_panes = "A4"
    return sheet, records


def build_override_sheet(workbook, template_sheet):
    sheet = workbook.create_sheet("人工营业额覆盖")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "人工营业额覆盖（优先级高于数据源；仅维护经确认的例外）"
    sheet["A1"].font = Font(size=15, bold=True, color="173D3A")
    for col, header in enumerate(OVERRIDE_HEADERS, start=1):
        sheet.cell(3, col, header)
    style_headers(sheet, 3, [14, 30, 18, 34, 14, 18, 36])
    records = []
    for row in range(3, template_sheet.max_row + 1):
        revenue_cell = template_sheet.cell(row, 11)
        if template_sheet.cell(row, 2).value in (None, "") or revenue_cell.data_type == "f":
            continue
        records.append([
            template_sheet.cell(row, 4).value, template_sheet.cell(row, 5).value,
            template_sheet.cell(row, 8).value, template_sheet.cell(row, 9).value,
            template_sheet.cell(row, 10).value, revenue_cell.value, "原模板人工营业额，已预置",
        ])
    # Keep a blank table row if a future template has no manual overrides.
    if not records:
        records = [[None] * len(OVERRIDE_HEADERS)]
    for output_row, values in enumerate(records, start=4):
        for col, value in enumerate(values, start=1):
            sheet.cell(output_row, col, value)
            sheet.cell(output_row, col).alignment = Alignment(vertical="top", wrap_text=True)
        sheet.cell(output_row, 6).number_format = '#,##0.00'
    add_table(sheet, "tblRevenueOverride", f"A3:G{3 + len(records)}")
    sheet.freeze_panes = "A4"
    return sheet


def output_formula():
    return (
        '=LET('
        'm,tblRevenueInput[档口唯一标志]<>"",'
        'sid,FILTER(tblRevenueInput[档口唯一标志],m,""),'
        'sbranch,FILTER(tblRevenueInput[分部名称],m,""),'
        'scode,FILTER(tblRevenueInput[项目代码],m,""),'
        'sname,FILTER(tblRevenueInput[项目名称],m,""),'
        'stype,FILTER(tblRevenueInput[项目类型],m,""),'
        'sgroup,FILTER(tblRevenueInput[分组作业名称],m,""),'
        'swindow,FILTER(tblRevenueInput[分组作业窗口名称],m,""),'
        'ssigner,FILTER(tblRevenueInput[签订人],m,""),'
        'samount,FILTER(tblRevenueInput[全额营业额合计],m,""),'
        'srow,FILTER(ROW(tblRevenueInput[档口唯一标志]),m,""),'
        'sk,scode&"|"&sname&"|"&sgroup&"|"&swindow&"|"&ssigner,'
        'sf,scode&"|"&sname&"|"&sgroup&"|"&swindow,'
        'bk,tblTargetBase[项目代码]&"|"&tblTargetBase[项目名称]&"|"&tblTargetBase[分组作业名称]&"|"&tblTargetBase[分组作业窗口名称]&"|"&tblTargetBase[责任人],'
        'bf,tblTargetBase[项目代码]&"|"&tblTargetBase[项目名称]&"|"&tblTargetBase[分组作业名称]&"|"&tblTargetBase[分组作业窗口名称],'
        'om,tblRevenueOverride[项目代码]<>"",'
        'ok,FILTER(tblRevenueOverride[项目代码]&"|"&tblRevenueOverride[项目名称]&"|"&tblRevenueOverride[分组作业名称]&"|"&tblRevenueOverride[分组作业窗口名称]&"|"&tblRevenueOverride[责任人],om,""),'
        'oa,FILTER(tblRevenueOverride[营业额覆盖],om,""),'
        'ov,IFERROR(XLOOKUP(bk,ok,oa,""),""),'
        'sr,IFERROR(XLOOKUP(bk,sk,srow,""),""),'
        'sa,IFERROR(XLOOKUP(bk,sk,samount,""),""),'
        'amount,IF(ov<>"",ov,sa),'
        'status,IF(ov<>"","保留人工录入",IF(sr<>"","正常","待核验")), '
        'mark,IF(tblTargetBase[档口唯一标志]<>"","原目标",""),'
        'note,IF(status="待核验","原稿未匹配",IF(status="保留人工录入","按人工覆盖值保留","五项键精确匹配")), '
        'baseOut,HSTACK(status,sr,mark,tblTargetBase[档口唯一标志],tblTargetBase[分部名称],tblTargetBase[项目代码],tblTargetBase[项目名称],tblTargetBase[项目类型],tblTargetBase[分组作业名称],tblTargetBase[分组作业窗口名称],tblTargetBase[责任人],amount,note),'
        'newMask,ISNA(XMATCH(sk,bk)), '
        'newStatus,IF(ISNUMBER(XMATCH(sf,bf)),"新增（责任人变更）","新增窗口"),'
        'newOut,IFERROR(HSTACK(FILTER(newStatus,newMask),FILTER(srow,newMask),FILTER(IF(sid<>"","新增",""),newMask),FILTER(sid,newMask),FILTER(sbranch,newMask),FILTER(scode,newMask),FILTER(sname,newMask),FILTER(stype,newMask),FILTER(sgroup,newMask),FILTER(swindow,newMask),FILTER(ssigner,newMask),FILTER(samount,newMask),FILTER(IF(newStatus="新增（责任人变更）","同窗口责任人变更，单独新增","原稿新增窗口"),newMask)),""),'
        'allRows,VSTACK(baseOut,newOut),FILTER(allRows,INDEX(allRows,,4)<>"")'
        ')'
    )


def build_output_sheet(workbook):
    sheet = workbook.create_sheet("函数目标输出")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "函数目标输出（无需手工填数）"
    sheet["A1"].font = Font(size=15, bold=True, color="173D3A")
    sheet["A2"] = "以下数据由“营业额数据源”和“人工营业额覆盖”实时计算；新增窗口与责任人变更自动追加。"
    sheet["A2"].alignment = Alignment(wrap_text=True)
    for col, header in enumerate(OUTPUT_HEADERS, start=1):
        sheet.cell(3, col, header)
    style_headers(sheet, 3, [16, 12, 10, 16, 18, 14, 30, 14, 18, 34, 14, 18, 34])
    sheet["A4"] = output_formula()
    sheet["L4"].number_format = '#,##0.00'
    sheet.freeze_panes = "A4"
    return sheet


def build_audit_sheet(workbook):
    sheet = workbook.create_sheet("营业额导入核验")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "营业额导入核验（纯函数）"
    sheet["A1"].font = Font(size=15, bold=True, color="173D3A")
    summary = [
        ("输出记录数", "=ROWS('函数目标输出'!A4#)"),
        ("正常", "=COUNTIF(INDEX('函数目标输出'!A4#,0,1),\"正常\")"),
        ("保留人工营业额", "=COUNTIF(INDEX('函数目标输出'!A4#,0,1),\"保留人工录入\")"),
        ("待核验", "=COUNTIF(INDEX('函数目标输出'!A4#,0,1),\"待核验\")"),
        ("新增窗口", "=COUNTIF(INDEX('函数目标输出'!A4#,0,1),\"新增窗口\")"),
        ("责任人变更", "=COUNTIF(INDEX('函数目标输出'!A4#,0,1),\"新增（责任人变更）\")"),
    ]
    for row, (label, formula) in enumerate(summary, start=3):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, formula)
    sheet["A10"] = "明细：点击“数据源行”跳转到本文件内的营业额数据源。"
    sheet["A10"].fill = HEADER_FILL
    for col, header in enumerate(OUTPUT_HEADERS, start=1):
        sheet.cell(12, col, "数据源行（点击跳转）" if col == 2 else header)
    style_headers(sheet, 12, [16, 16, 10, 16, 18, 14, 30, 14, 18, 34, 14, 18, 34])
    sheet["A13"] = (
        '=LET(x,\'函数目标输出\'!A4#,'
        'HSTACK(INDEX(x,,1),IF(INDEX(x,,2)=\"\",\"\",HYPERLINK("#\'营业额数据源\'!A"&INDEX(x,,2),INDEX(x,,2))),DROP(x,,3)))'
    )
    sheet["L13"].number_format = '#,##0.00'
    sheet.freeze_panes = "A13"
    sheet.conditional_formatting.add("A13:M1200", FormulaRule(formula=['$A13="待核验"'], fill=WARN_FILL))
    sheet.conditional_formatting.add("A13:M1200", FormulaRule(formula=['LEFT($A13,2)="新增"'], fill=NEW_FILL))
    sheet.conditional_formatting.add("A13:M1200", FormulaRule(formula=['$A13="保留人工录入"'], fill=MANUAL_FILL))
    return sheet


def build(source_path: Path, template_path: Path, output_path: Path):
    workbook = load_workbook(template_path, data_only=False, keep_links=True)
    source_workbook = load_workbook(source_path, data_only=False, keep_links=True)
    source_sheet = source_workbook.worksheets[0]
    template_sheet = find_sheet(workbook, "(8月)")
    clear_generated_sheets(workbook)
    build_instructions(workbook)
    build_input_sheet(workbook, source_sheet)
    build_base_sheet(workbook, template_sheet)
    build_override_sheet(workbook, template_sheet)
    build_output_sheet(workbook)
    build_audit_sheet(workbook)
    final_title = prepare_final_sheet(workbook, template_sheet)

    # Keep daily use to five clear tabs.  The two calculation helpers and the
    # historical cost-reference sheets stay in the workbook but do not distract.
    for title in ["目标基准", "函数目标输出"]:
        workbook[title].sheet_state = "hidden"
    for sheet in workbook.worksheets:
        if sheet.title.startswith(" 2026年") or sheet.title == "Sheet1":
            sheet.sheet_state = "hidden"
    for title in ["使用说明", final_title, "营业额数据源", "营业额导入核验", "人工营业额覆盖"]:
        workbook[title].sheet_state = "visible"
    visible_order = ["使用说明", final_title, "营业额数据源", "营业额导入核验", "人工营业额覆盖"]
    hidden_order = [sheet.title for sheet in workbook.worksheets if sheet.title not in visible_order]
    workbook._sheets = [workbook[title] for title in visible_order + hidden_order]
    workbook.active = 1
    # All revenue formulas now point to in-workbook tables.  Remove the old
    # external-link package so Excel/WPS no longer raises a refresh warning.
    workbook._external_links = []
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="生成营业额自动导入纯函数 Excel 模板。")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.template, args.output)
    print(f"已生成：{args.output}")


if __name__ == "__main__":
    main()
